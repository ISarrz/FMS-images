"""
FMS-images — настоящие PNG-скриншоты расписания уроков из Excel.

Каждый лист книги (.xlsx) — это отдельная параллель (например, «10», «11»).
Скрипт рендерит сами листы Excel «как есть» (с родным оформлением: цвета,
границы, шрифты) и сохраняет по одному PNG на параллель в каталог output/.

Файл НЕ перерисовывается вручную — рендеринг делает LibreOffice:
    .xlsx  ->  PDF (LibreOffice, каждый лист на отдельной странице)
           ->  PNG (pdftoppm, по странице на лист)
           ->  обрезка белых полей (Pillow)

Использование:
    python generate.py                      # обработать все .xlsx из input/
    python generate.py input/05.09.2026.xlsx
    python generate.py file.xlsx --input-dir input --output-dir output --dpi 200

Требуются системные пакеты: libreoffice (soffice) и poppler-utils (pdftoppm).
"""
from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from typing import List, Optional

import openpyxl
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image, ImageChops

CROP_PADDING = 12       # белое поле, оставляемое вокруг таблицы после обрезки, px
SOFFICE_TIMEOUT = 180   # таймаут конвертации в PDF, сек


def _find_binary(names: List[str]) -> Optional[str]:
    for name in names:
        path = shutil.which(name)
        if path:
            return path
    return None


def sheet_is_empty(sheet) -> bool:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                return False
    return True


def prepare_workbook(src_path: str, tmp_dir: str) -> tuple[str, List[str]]:
    """Готовит копию книги к печати: каждый непустой лист — на одну страницу.

    Возвращает путь к временному .xlsx и список имён листов в том же порядке,
    в котором они окажутся страницами PDF.
    """
    workbook = openpyxl.load_workbook(src_path)

    kept: List[str] = []
    for name in list(workbook.sheetnames):
        sheet = workbook[name]
        if sheet_is_empty(sheet):
            del workbook[name]
            continue
        kept.append(name)

        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_margins = PageMargins(
            left=0.2, right=0.2, top=0.2, bottom=0.2, header=0, footer=0
        )
        if sheet.dimensions:
            sheet.print_area = sheet.dimensions

    tmp_xlsx = os.path.join(tmp_dir, "fit.xlsx")
    workbook.save(tmp_xlsx)
    return tmp_xlsx, kept


def xlsx_to_pdf(soffice: str, xlsx_path: str, tmp_dir: str) -> str:
    """Конвертирует .xlsx в PDF через LibreOffice в headless-режиме."""
    profile_uri = "file://" + os.path.join(tmp_dir, "lo_profile")
    cmd = [
        soffice,
        "--headless",
        "--norestore",
        "--convert-to", "pdf",
        "--outdir", tmp_dir,
        "-env:UserInstallation=" + profile_uri,
        xlsx_path,
    ]
    subprocess.run(
        cmd, check=True, timeout=SOFFICE_TIMEOUT,
        stdout=subprocess.PIPE, stderr=subprocess.PIPE,
    )
    pdf_path = os.path.join(tmp_dir, os.path.splitext(os.path.basename(xlsx_path))[0] + ".pdf")
    if not os.path.isfile(pdf_path):
        raise RuntimeError(f"LibreOffice не создал PDF: {pdf_path}")
    return pdf_path


def pdf_to_pngs(pdftoppm: str, pdf_path: str, tmp_dir: str, dpi: int) -> List[str]:
    """Разбивает PDF на PNG постранично (одна страница = один лист)."""
    prefix = os.path.join(tmp_dir, "page")
    subprocess.run(
        [pdftoppm, "-png", "-r", str(dpi), pdf_path, prefix],
        check=True, timeout=SOFFICE_TIMEOUT,
        stdout=subprocess.PIPE, stderr=subprocess.PIPE,
    )
    return sorted(
        os.path.join(tmp_dir, f)
        for f in os.listdir(tmp_dir)
        if f.startswith("page") and f.endswith(".png")
    )


def autocrop(png_path: str, out_path: str) -> None:
    """Обрезает белые поля вокруг таблицы и сохраняет результат."""
    image = Image.open(png_path).convert("RGB")
    background = Image.new("RGB", image.size, (255, 255, 255))
    bbox = ImageChops.difference(image, background).getbbox()
    if bbox:
        left, top, right, bottom = bbox
        left = max(0, left - CROP_PADDING)
        top = max(0, top - CROP_PADDING)
        right = min(image.width, right + CROP_PADDING)
        bottom = min(image.height, bottom + CROP_PADDING)
        image = image.crop((left, top, right, bottom))
    image.save(out_path, format="PNG")


def process_workbook(src_path: str, output_dir: str, soffice: str, pdftoppm: str, dpi: int) -> List[str]:
    """Обрабатывает книгу: по одному PNG-скриншоту на лист-параллель."""
    base = os.path.splitext(os.path.basename(src_path))[0]
    os.makedirs(output_dir, exist_ok=True)
    saved: List[str] = []

    with tempfile.TemporaryDirectory(prefix="fms-images-") as tmp_dir:
        tmp_xlsx, sheet_names = prepare_workbook(src_path, tmp_dir)
        if not sheet_names:
            print("  нет непустых листов — пропущено")
            return saved

        pdf_path = xlsx_to_pdf(soffice, tmp_xlsx, tmp_dir)
        pages = pdf_to_pngs(pdftoppm, pdf_path, tmp_dir, dpi)

        if len(pages) != len(sheet_names):
            print(
                f"  предупреждение: страниц PDF ({len(pages)}) не совпадает с числом "
                f"листов ({len(sheet_names)}); сопоставляю по порядку"
            )

        for page_png, sheet_name in zip(pages, sheet_names):
            safe = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in sheet_name)
            out_path = os.path.join(output_dir, f"{base}_{safe}.png")
            autocrop(page_png, out_path)
            saved.append(out_path)
            print(f"  параллель {sheet_name!r:>6} -> {out_path}")

    return saved


def collect_inputs(args) -> List[str]:
    if args.files:
        return args.files
    if not os.path.isdir(args.input_dir):
        return []
    return sorted(
        os.path.join(args.input_dir, name)
        for name in os.listdir(args.input_dir)
        if name.lower().endswith(".xlsx") and not name.startswith("~$")
    )


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Настоящие PNG-скриншоты расписания из Excel по параллелям (листам)."
    )
    parser.add_argument("files", nargs="*", help="конкретные .xlsx (по умолчанию — все из --input-dir)")
    parser.add_argument("--input-dir", default="input", help="каталог с .xlsx (по умолчанию: input)")
    parser.add_argument("--output-dir", default="output", help="куда сохранять PNG (по умолчанию: output)")
    parser.add_argument("--dpi", type=int, default=200, help="разрешение рендеринга (по умолчанию: 200)")
    args = parser.parse_args()

    soffice = _find_binary(["libreoffice", "soffice"])
    pdftoppm = _find_binary(["pdftoppm"])
    missing = []
    if not soffice:
        missing.append("libreoffice (soffice)")
    if not pdftoppm:
        missing.append("poppler-utils (pdftoppm)")
    if missing:
        print("Не найдены системные зависимости: " + ", ".join(missing))
        print("Установите их, например: sudo dnf install libreoffice-calc poppler-utils")
        return 1

    inputs = collect_inputs(args)
    if not inputs:
        print("Не найдено ни одного .xlsx. Положите файлы в каталог input/ или укажите путь явно.")
        return 1

    total = 0
    for xlsx_path in inputs:
        if not os.path.isfile(xlsx_path):
            print(f"Файл не найден: {xlsx_path}")
            continue
        print(f"Обрабатываю {xlsx_path}")
        try:
            total += len(process_workbook(xlsx_path, args.output_dir, soffice, pdftoppm, args.dpi))
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, RuntimeError) as exc:
            print(f"  ошибка при обработке {xlsx_path}: {exc}")

    print(f"Готово. Сохранено изображений: {total}")
    return 0 if total else 1


if __name__ == "__main__":
    sys.exit(main())
