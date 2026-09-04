"""Рендеринг листов Excel в PNG-скриншоты.

Конвейер: .xlsx -> PDF (LibreOffice, лист = страница) -> PNG (pdftoppm) ->
обрезка белых полей (Pillow). Листы не перерисовываются — сохраняется родное
оформление Excel.
"""
from __future__ import annotations

import logging
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import List, Optional

import openpyxl
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties
from PIL import Image, ImageChops

from .errors import DependencyError, RenderError

logger = logging.getLogger(__name__)

DEFAULT_DPI = 200
DEFAULT_CROP_PADDING = 12
DEFAULT_TIMEOUT = 180


@dataclass
class RenderTools:
    """Пути к внешним программам, нужным для рендеринга."""

    soffice: str
    pdftoppm: str

    @staticmethod
    def _which(names: List[str]) -> Optional[str]:
        for name in names:
            path = shutil.which(name)
            if path:
                return path
        return None

    @classmethod
    def discover(cls) -> "RenderTools":
        """Находит libreoffice/soffice и pdftoppm в PATH или бросает :class:`DependencyError`."""
        soffice = cls._which(["libreoffice", "soffice"])
        pdftoppm = cls._which(["pdftoppm"])
        missing = []
        if not soffice:
            missing.append("libreoffice (soffice)")
        if not pdftoppm:
            missing.append("poppler-utils (pdftoppm)")
        if missing:
            raise DependencyError(
                "Не найдены системные зависимости: " + ", ".join(missing)
                + ". Установите их, например: sudo dnf install libreoffice-calc poppler-utils"
            )
        return cls(soffice=soffice, pdftoppm=pdftoppm)


def _sheet_is_empty(sheet) -> bool:
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                return False
    return True


def _safe_name(name: str) -> str:
    return "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in name.strip())


def _prepare_workbook(src_path: str, tmp_dir: str) -> tuple[str, List[str]]:
    """Готовит копию книги к печати (лист = одна страница). Возвращает путь и список листов."""
    workbook = openpyxl.load_workbook(src_path)

    kept: List[str] = []
    for name in list(workbook.sheetnames):
        sheet = workbook[name]
        if _sheet_is_empty(sheet):
            del workbook[name]
            continue
        kept.append(name)

        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 1
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_margins = PageMargins(left=0.2, right=0.2, top=0.2, bottom=0.2, header=0, footer=0)
        if sheet.dimensions:
            sheet.print_area = sheet.dimensions

    tmp_xlsx = os.path.join(tmp_dir, "fit.xlsx")
    workbook.save(tmp_xlsx)
    return tmp_xlsx, kept


def _xlsx_to_pdf(tools: RenderTools, xlsx_path: str, tmp_dir: str, timeout: int) -> str:
    profile_uri = "file://" + os.path.join(tmp_dir, "lo_profile")
    cmd = [
        tools.soffice, "--headless", "--norestore",
        "--convert-to", "pdf", "--outdir", tmp_dir,
        "-env:UserInstallation=" + profile_uri, xlsx_path,
    ]
    try:
        subprocess.run(cmd, check=True, timeout=timeout,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        raise RenderError(f"LibreOffice не смог конвертировать {xlsx_path}: {exc}") from exc

    pdf_path = os.path.join(tmp_dir, Path(xlsx_path).stem + ".pdf")
    if not os.path.isfile(pdf_path):
        raise RenderError(f"LibreOffice не создал PDF: {pdf_path}")
    return pdf_path


def _pdf_to_pngs(tools: RenderTools, pdf_path: str, tmp_dir: str, dpi: int, timeout: int) -> List[str]:
    prefix = os.path.join(tmp_dir, "page")
    try:
        subprocess.run([tools.pdftoppm, "-png", "-r", str(dpi), pdf_path, prefix],
                       check=True, timeout=timeout,
                       stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as exc:
        raise RenderError(f"pdftoppm не смог разобрать PDF {pdf_path}: {exc}") from exc

    return sorted(
        os.path.join(tmp_dir, f)
        for f in os.listdir(tmp_dir)
        if f.startswith("page") and f.endswith(".png")
    )


def _autocrop(src_png: str, dst_png: str, padding: int) -> None:
    image = Image.open(src_png).convert("RGB")
    background = Image.new("RGB", image.size, (255, 255, 255))
    bbox = ImageChops.difference(image, background).getbbox()
    if bbox:
        left, top, right, bottom = bbox
        left = max(0, left - padding)
        top = max(0, top - padding)
        right = min(image.width, right + padding)
        bottom = min(image.height, bottom + padding)
        image = image.crop((left, top, right, bottom))
    image.save(dst_png, format="PNG")


def render_workbook(
    src_path: str,
    output_dir: str,
    *,
    dpi: int = DEFAULT_DPI,
    tools: Optional[RenderTools] = None,
    crop_padding: int = DEFAULT_CROP_PADDING,
    timeout: int = DEFAULT_TIMEOUT,
) -> List[Path]:
    """Рендерит книгу Excel в PNG — по одному на лист-параллель.

    Все изображения одной книги складываются в подпапку ``output_dir/<имя_книги>/``,
    файлы называются по имени листа: ``<имя_листа>.png``.

    :returns: список путей к созданным изображениям (может быть пустым).
    """
    tools = tools or RenderTools.discover()
    base = Path(src_path).stem
    book_dir = Path(output_dir) / base
    saved: List[Path] = []

    with tempfile.TemporaryDirectory(prefix="fms-images-") as tmp_dir:
        tmp_xlsx, sheet_names = _prepare_workbook(src_path, tmp_dir)
        if not sheet_names:
            logger.info("%s: нет непустых листов — пропущено", src_path)
            return saved

        pdf_path = _xlsx_to_pdf(tools, tmp_xlsx, tmp_dir, timeout)
        pages = _pdf_to_pngs(tools, pdf_path, tmp_dir, dpi, timeout)

        if len(pages) != len(sheet_names):
            logger.warning(
                "%s: страниц PDF (%d) не совпадает с числом листов (%d); сопоставляю по порядку",
                src_path, len(pages), len(sheet_names),
            )

        book_dir.mkdir(parents=True, exist_ok=True)
        for page_png, sheet_name in zip(pages, sheet_names):
            out_path = book_dir / f"{_safe_name(sheet_name)}.png"
            _autocrop(page_png, str(out_path), crop_padding)
            saved.append(out_path)
            logger.info("параллель %r -> %s", sheet_name, out_path)

    return saved


def iter_workbooks(input_dir: str) -> List[Path]:
    """Список .xlsx в каталоге (без временных файлов Excel)."""
    directory = Path(input_dir)
    if not directory.is_dir():
        return []
    return sorted(
        p for p in directory.iterdir()
        if p.suffix.lower() == ".xlsx" and not p.name.startswith("~$")
    )


def render_directory(
    input_dir: str,
    output_dir: str,
    *,
    dpi: int = DEFAULT_DPI,
    tools: Optional[RenderTools] = None,
    crop_padding: int = DEFAULT_CROP_PADDING,
    timeout: int = DEFAULT_TIMEOUT,
) -> List[Path]:
    """Рендерит все .xlsx из каталога. Возвращает пути ко всем созданным изображениям."""
    tools = tools or RenderTools.discover()
    saved: List[Path] = []
    for xlsx_path in iter_workbooks(input_dir):
        logger.info("обрабатываю %s", xlsx_path)
        saved.extend(render_workbook(
            str(xlsx_path), output_dir,
            dpi=dpi, tools=tools, crop_padding=crop_padding, timeout=timeout,
        ))
    return saved
